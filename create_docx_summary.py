from docx import Document

if __name__ == '__main__':
    doc = Document()
    doc.add_heading('Gastric Cancer Histopathology Classification Project', level=1)

    doc.add_heading('1) Project Overview', level=2)
    doc.add_paragraph(
        'This project implements a medical image classification framework for multi-class gastric cancer histopathology. '
        'It uses PyTorch with a ConvNeXt ensemble, Grad-CAM interpretability, and comprehensive training and evaluation tooling. '
        'The target classes are ADI, DEB, LYM, and MUC.'
    )

    doc.add_heading('2) Model Architecture', level=2)
    doc.add_paragraph(
        'The core model is a ConvNeXt ensemble composed of ConvNeXt-Tiny, ConvNeXt-Small, and ConvNeXt-Base backbones. '
        'Each backbone replaces the default classification head with a dropout layer and a linear output layer matching the four classes. '
        'The ensemble output is computed by stacking each model output and applying learnable weights for weighted averaging.'
    )

    doc.add_heading('3) Model Factory & Vit alternatives', level=2)
    doc.add_paragraph(
        'The project exposes a model factory function to build either a single ConvNeXt classifier or the full ensemble. '
        'Supported model names include convnext_tiny, convnext_small, convnext_base, and convnext_ensemble. '
        'A Vision Transformer (ViT) alternative is not currently implemented in this repository, but a natural future extension would be to add ViT backbone support via torchvision or timm, allowing direct comparison between ConvNeXt and transformer architectures.'
    )

    doc.add_heading('4) Training Configuration', level=2)
    doc.add_paragraph('Training is configured in config.yaml. Key settings include:')
    doc.add_paragraph('• num_classes: 4', style='List Bullet')
    doc.add_paragraph('• image_size: 224', style='List Bullet')
    doc.add_paragraph('• batch_size: 32', style='List Bullet')
    doc.add_paragraph('• model ensemble: convnext_tiny, convnext_small, convnext_base', style='List Bullet')
    doc.add_paragraph('• optimizer: AdamW, lr=0.0001, weight_decay=0.0001', style='List Bullet')
    doc.add_paragraph('• scheduler: CosineAnnealingLR', style='List Bullet')
    doc.add_paragraph('• dropout: 0.2', style='List Bullet')
    doc.add_paragraph('• epochs: 20', style='List Bullet')
    doc.add_paragraph('• early stopping patience: 8', style='List Bullet')

    doc.add_heading('5) Evaluation Results', level=2)
    doc.add_paragraph(
        'Evaluation generates detailed metrics including accuracy, precision, recall, macro F1-score, log loss, and multi-class ROC-AUC. '
        'The repository includes reports such as test_metrics.json, confusion matrices, per-class metrics plots, and ROC curve visualizations. '
        'Grad-CAM visualizations are also produced for model interpretability on test images.'
    )

    doc.add_heading('6) Interpretation & Insights', level=2)
    doc.add_paragraph(
        'Interpretability is provided through Grad-CAM attention maps, which highlight regions contributing most to each prediction. '
        'This allows practitioners to verify that the ensemble is focusing on relevant histopathological structures rather than background noise. '
        'The ensemble approach improves robustness by combining the strengths of lightweight, balanced, and higher-capacity ConvNeXt variants.'
    )

    doc.add_paragraph(
        'Overall, the project is built as a practical research pipeline with configurable training, evaluation, visualization, and inference components suitable for gastric cancer histopathology classification tasks.'
    )

    output_file = 'Gastric_Cancer_Classification_Project_Summary.docx'
    doc.save(output_file)
    print(f'Created {output_file}')
